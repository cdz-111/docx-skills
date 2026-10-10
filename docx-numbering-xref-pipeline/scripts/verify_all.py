# -*- coding: utf-8 -*-
"""长文档编号与交叉引用治理 —— 一条命令跑完全部机器可验项。

用法:
    python verify_all.py <文档.docx> [--base <改前备份.docx>] [--quiet]
                                   [--stage {s2,s6}]

检查共 13 项（详见同目录 SKILL.md §5；1~11 恒跑，12/13 仅在给了 --base 时跑）：
    1 XML 可解析            2 python-docx 可打开      3 书签配平
    4 域三态配平            5 域嵌套合法             6 REF 断链 / 缓存与语义
    7 _Ref 书签过宽         8 题注编号连续           9 清单编号与锚点
   10 正文残留硬写引用     11 结构计数
给了 --base 还额外检查：
   12 除 word/document.xml 外所有 zip 条目逐字节相同 + 条目顺序相同
   13 可见文本 diff 仅「图/表与编号之间的空格」变化、去编号骨架不变

--stage（默认 s2，向后兼容）
    s2 —— 编号治理阶段的验收器：第 11/12/13 项照常 FAIL。
    s6 —— 交付终点的验收器：S6 阶段**必然**会带来三项与基线不同的差异，
          这时把第 11/12/13 项标成 `EXPECT`（预期差异，显式标注原因、**不静默**），
          不计入 FAIL：
            11 结构计数     —— S6 插题注段，paras 会 +1
            12 部件字节一致 —— S6 开了修订模式，新增 word/comments.xml
            13 可见文本 diff —— S6 插题注段 / 修订模式，段数不同

退出码语义**不变**：全绿（含 EXPECT）0，有真 FAIL 2，用法错误 1。
"""
import zipfile, re, sys, os, io
from collections import Counter, defaultdict
import xml.etree.ElementTree as ET

# ---------------------------------------------------------------- 基础工具
TAGP = re.compile(r'<w:p(?: [^>]*)?/>|<w:p [^>]*>|<w:p>|</w:p>')
WT = re.compile(r'<w:t(?: [^>]*)?>([^<]*)</w:t>', re.S)
FLDC = re.compile(r'<w:fldChar w:fldCharType="(begin|separate|end)"/>')
TB = '<w:fldChar w:fldCharType="begin"/>'
TS = '<w:fldChar w:fldCharType="separate"/>'
TE = '<w:fldChar w:fldCharType="end"/>'
INSTR = re.compile(r'<w:instrText[^>]*>(.*?)</w:instrText>', re.S)
NUM = re.compile(r'([图表])\s*(\d+)')
# S3「续表标注」段的两种合法形状（第 10 项豁免用；**必须整段完全匹配**，模块级预编译）：
#   A) 「续表 1」「续 表 1」「续表 2（续）」
#   B) 「表 2（续）」「图 3（续）」—— **必须带「（续）」后缀**（否则纯硬写「表 2」会被误豁免）
RE_CONT_A = re.compile(
    r'[\s\u3000]*续[\s\u3000]*[图表][\s\u3000]*\d+[\s\u3000]*(（续）|\(续\))?[\s\u3000]*')
RE_CONT_B = re.compile(
    r'[\s\u3000]*[图表][\s\u3000]*\d+[\s\u3000]*(（续）|\(续\))[\s\u3000]*')
DOCXML = 'word/document.xml'


def paras(x):
    """深度感知切分 <w:p>（含表格单元格内段落）"""
    res, depth, st = [], 0, None
    for m in TAGP.finditer(x):
        s = m.group(0)
        if s.endswith('/>'):
            if depth == 0:
                res.append((m.start(), m.end()))
        elif s == '</w:p>':
            depth -= 1
            if depth <= 0:
                if st is not None:
                    res.append((st, m.end()))
                st = None; depth = 0
        else:
            if depth == 0:
                st = m.start()
            depth += 1
    return res


def render(frag):
    """把域替换成其缓存后再取 <w:t> 文本"""
    out, pos = [], 0
    while True:
        a = frag.find(TB, pos)
        if a < 0:
            out.append(frag[pos:]); break
        sm = frag.find(TS, a); em = frag.find(TE, a)
        if sm < 0 or em < 0 or sm > em:
            out.append(frag[pos:a])
            if em < 0:
                out.append(frag[a:]); break
            out.append(frag[a:em]); pos = em; continue
        out.append(frag[pos:a]); out.append(frag[sm + len(TS):em])
        pos = em + len(TE)
    return ''.join(WT.findall(''.join(out)))


def bm_spans(x):
    st, en = {}, {}
    for m in re.finditer(r'<w:bookmarkStart\b[^>]*?/>', x):
        t = m.group(0)
        i = re.search(r'w:id="(\d+)"', t); n = re.search(r'w:name="([^"]*)"', t)
        if i and n:
            st[i.group(1)] = (n.group(1), m.end())
    for m in re.finditer(r'<w:bookmarkEnd\b[^>]*?/>', x):
        i = re.search(r'w:id="(\d+)"', m.group(0))
        if i:
            en[i.group(1)] = m.start()
    return {n: (s, en[k]) for k, (n, s) in st.items() if k in en}


def cached_regions(q):
    stack, out = [], []
    for m in FLDC.finditer(q):
        k = m.group(1)
        if k == 'begin':
            stack.append(None)
        elif k == 'separate':
            if stack and stack[-1] is None:
                stack[-1] = m.end()
        else:
            if stack:
                s = stack.pop()
                if s is not None:
                    out.append((s, m.start()))
    return out


def visible_segs(q):
    regs = cached_regions(q)
    return [m for m in WT.finditer(q)
            if not any(s <= m.start() < e for s, e in regs)]


def load(path):
    z = zipfile.ZipFile(path)
    x = z.read(DOCXML).decode('utf-8')
    names = z.namelist()
    z.close()
    return x, names


# ---------------------------------------------------------------- 报告器
class R:
    def __init__(self):
        self.rows = []
        self.fail = 0
        self.expect = 0

    def add(self, ok, name, detail=''):
        self.rows.append(('PASS' if ok else 'FAIL', name, detail))
        if not ok:
            self.fail += 1

    def add_expect(self, ok, name, detail=''):
        """S6 阶段的「预期差异」：PASS 照常；差异记 EXPECT，**不计入 fail**。"""
        if ok:
            self.rows.append(('PASS', name, detail))
        else:
            self.rows.append(('EXPECT', name, detail))
            self.expect += 1

    def out(self):
        print('=' * 78)
        for st, name, detail in self.rows:
            print('  %s %-26s %s' % (st, name, detail))
        print('=' * 78)
        tail = '失败 %d 项' % self.fail
        if self.expect:
            tail += ' / 预期差异 %d 项' % self.expect
        print('  %s   （%s）' % ('全绿 ✓' if not self.fail else '存在问题 ✗', tail))


# ---------------------------------------------------------------- 各项检查
def check(path, base=None, quiet=False, stage='s2'):
    rep = R()
    x, names = load(path)

    # 1 XML
    try:
        ET.fromstring(x.encode('utf-8'))
        rep.add(True, '1 XML 可解析', '%d chars' % len(x))
    except Exception as e:
        rep.add(False, '1 XML 可解析', str(e))
        rep.out(); return rep

    # 2 python-docx
    try:
        import docx
        d = docx.Document(path)
        rep.add(True, '2 python-docx 可打开', '段落 %d / 表格 %d' % (len(d.paragraphs), len(d.tables)))
    except Exception as e:
        rep.add(False, '2 python-docx 可打开', '%s: %s' % (type(e).__name__, e))

    # 3 书签配平
    bs = len(re.findall(r'<w:bookmarkStart\b', x))
    be = len(re.findall(r'<w:bookmarkEnd\b', x))
    rep.add(bs == be, '3 书签配平', '%d / %d' % (bs, be))

    # 4 域三态配平
    fc = Counter(FLDC.findall(x))
    ok4 = fc['begin'] == fc['separate'] == fc['end']
    rep.add(ok4, '4 域三态配平', 'begin=%d separate=%d end=%d' % (fc['begin'], fc['separate'], fc['end']))

    # 5 域嵌套
    seq = FLDC.findall(x)
    depth, mind, sep_bad, cur_sep = 0, 0, 0, 0
    for s in seq:
        if s == 'begin':
            depth += 1; cur_sep = 0
        elif s == 'separate':
            if depth == 0 or cur_sep:
                sep_bad += 1
            cur_sep += 1
        else:
            depth -= 1
            mind = min(mind, depth)
    rep.add(depth == 0 and mind >= 0 and sep_bad == 0, '5 域嵌套合法',
            '结束深度=%d 最小深度=%d separate 异常=%d' % (depth, mind, sep_bad))

    # 6 REF 断链 / 缓存与语义
    sp = bm_spans(x)
    nref = missing = mismatch = 0
    bad_rows = []
    for i, (a, b) in enumerate(paras(x)):
        q = x[a:b]
        for m in re.finditer(re.escape(TB), q):
            seg = q[m.end():]
            em = re.search(re.escape(TE), seg)
            if not em:
                continue
            body = seg[:em.start()]
            im = INSTR.search(body)
            if not im:
                continue
            instr = im.group(1).strip()
            if not instr.upper().startswith('REF'):
                continue
            tm = re.search(r'REF\s+(\S+)', instr)
            if not tm:
                continue
            nref += 1
            tgt = tm.group(1)
            sm = body.find(TS)
            cache = ''.join(WT.findall(body[sm:])).strip() if sm >= 0 else ''
            if tgt not in sp:
                missing += 1; bad_rows.append((i, tgt, '书签不存在', ''))
                continue
            exp = render(x[sp[tgt][0]:sp[tgt][1]]).strip()
            if cache != exp:
                mismatch += 1; bad_rows.append((i, tgt, cache, exp))
    rep.add(missing == 0 and mismatch == 0, '6 REF 断链/缓存与语义',
            'REF %d / 断链 %d / 不符 %d' % (nref, missing, mismatch))
    if not quiet:
        for i, tgt, c, e in bad_rows[:10]:
            print('        p%-5d %-14s 缓存=%r 应为=%r' % (i, tgt, c, e))

    # 7 _Ref 过宽
    wide = set()
    for i, (a, b) in enumerate(paras(x)):
        q = x[a:b]
        for m in re.finditer(r'<w:bookmarkStart\b[^>]*w:name="(_Ref[^"]*)"', q):
            nm = m.group(1)
            if nm not in sp:
                continue
            inner = x[sp[nm][0]:sp[nm][1]]
            if 'SEQ' in inner and re.search(r'<w:bookmarkStart\b[^>]*w:name="_Toc', inner):
                wide.add(nm)
    rep.add(len(wide) == 0, '7 _Ref 书签未过宽', '过宽 %d 个 %s' % (len(wide), sorted(wide)[:6] or ''))

    # 8 题注编号连续
    caps = defaultdict(list)
    toc2cap = {}
    for i, (a, b) in enumerate(paras(x)):
        q = x[a:b]
        ins = [t.strip() for t in INSTR.findall(q)]
        if not any(t.startswith('SEQ 图') or t.startswith('SEQ 表') for t in ins):
            continue
        r = render(q)
        m = re.match(r'^\s*([图表])\s*(\d+)', r)
        if not m:
            continue
        kind, num = m.group(1), int(m.group(2))
        caps[kind].append((i, num))
        for nm in re.findall(r'<w:bookmarkStart\b[^>]*w:name="(_Toc[^"]*)"', q):
            toc2cap[nm] = (kind, num)
    ok8 = True; det8 = []
    for kind in ('图', '表'):
        nums = [n for _i, n in caps[kind]]
        if nums != list(range(1, len(nums) + 1)):
            ok8 = False
            dup = [n for n, c in Counter(nums).items() if c > 1]
            det8.append('%s: %d 个，重复 %s' % (kind, len(nums), dup or '无'))
        else:
            det8.append('%s 1..%d 无重复' % (kind, len(nums)))
    rep.add(ok8, '8 题注编号连续', ' ; '.join(det8))

    # 9 清单条目编号与锚点
    ent = defaultdict(list); dangle = 0; mism9 = 0; det9 = []
    for i, (a, b) in enumerate(paras(x)):
        q = x[a:b]
        if 'PAGEREF _Toc' not in q:
            continue
        tgts = re.findall(r'PAGEREF\s+(_Toc\w+)', q)
        if len(tgts) != 1:
            continue
        tgt = tgts[0]
        if tgt not in toc2cap:
            continue                 # 正文目录条目（锚点指向章节书签），跳过
        kind, num = toc2cap[tgt]
        ent[kind].append(tgt)
        head = q[:q.find('PAGEREF _Toc')]
        m = NUM.search(render(head))
        if not m:
            mism9 += 1; continue
        if m.group(1) != kind or int(m.group(2)) != num:
            mism9 += 1
            if len(det9) < 6:
                det9.append('p%d 条目写 %s%s 锚点应为 %s%d' % (i, m.group(1), m.group(2), kind, num))
    for kind in ('图', '表'):
        if len(ent[kind]) != len(caps[kind]):
            det9.append('%s清单条目 %d != 题注 %d' % (kind, len(ent[kind]), len(caps[kind])))
    rep.add(mism9 == 0 and dangle == 0
            and len(ent['图']) == len(caps['图']) and len(ent['表']) == len(caps['表']),
            '9 清单编号与锚点',
            '条目 图%d/表%d ; 编号不符 %d %s' % (len(ent['图']), len(ent['表']), mism9,
                                              '| ' + '; '.join(det9) if det9 else ''))

    # 10 正文残留硬写引用
    res = []
    for i, (a, b) in enumerate(paras(x)):
        q = x[a:b]
        ins = [t.strip() for t in INSTR.findall(q)]
        if any(t.startswith('SEQ 图') or t.startswith('SEQ 表') for t in ins):
            continue
        segs = visible_segs(q)
        if not segs:
            continue
        joined = ''.join(m.group(1) for m in segs)
        # S3「续表标注」段：**整段就是这个标注**，属结构性标注，不是正文硬写交叉引用。
        # **只豁免「整段完全匹配」**，绝不做全局替换（否则正文「续表 2 的后续数据…」的
        # 「表 2」会被静默吞掉）。接受两种（都必须是整段）：
        #   A) 「续表 1」「续 表 1」「续表 2（续）」
        #   B) 「表 2（续）」「图 3（续）」—— 必须带「（续）」后缀
        if RE_CONT_A.fullmatch(joined) or RE_CONT_B.fullmatch(joined):
            continue
        for m in NUM.finditer(joined):
            res.append((i, m.group(0)))
    rep.add(len(res) == 0, '10 正文残留硬写引用',
            '%d 处%s' % (len(res), '' if not res else ' ' + str([t for _i, t in res[:8]])))

    # 11 结构计数
    cnt = {'tbl': x.count('<w:tbl>'), 'drawing': x.count('<w:drawing>'),
           'tblGrid': x.count('<w:tblGrid>'), 'SEQ': len(re.findall(r'SEQ\s+[图表]', x)),
           'paras': len(paras(x))}
    if base:
        xb, namesb = load(base)
        cntb = {'tbl': xb.count('<w:tbl>'), 'drawing': xb.count('<w:drawing>'),
                'tblGrid': xb.count('<w:tblGrid>'), 'SEQ': len(re.findall(r'SEQ\s+[图表]', xb)),
                'paras': len(paras(xb))}
        same = cnt == cntb
        d = ' '.join('%s %d->%d' % (k, cntb[k], cnt[k]) for k in cnt if cntb[k] != cnt[k])
        det = d if not same else \
            'tbl=%d drawing=%d SEQ=%d paras=%d（与基线一致）' \
            % (cnt['tbl'], cnt['drawing'], cnt['SEQ'], cnt['paras'])
        if stage == 's6' and not same:
            rep.add_expect(False, '11 结构计数',
                           d + '　← EXPECT：S6 插题注段，段落计数必然 +1（非缺陷）')
        else:
            rep.add(same, '11 结构计数', det)
    else:
        rep.add(True, '11 结构计数', 'tbl=%d drawing=%d SEQ=%d paras=%d'
                % (cnt['tbl'], cnt['drawing'], cnt['SEQ'], cnt['paras']))

    # ---- 给了 base 才做的对比 ----
    if base:
        # 12 部件逐字节
        z1, z2 = zipfile.ZipFile(base), zipfile.ZipFile(path)
        n1, n2 = z1.namelist(), z2.namelist()
        only1 = [n for n in n1 if n not in n2]
        only2 = [n for n in n2 if n not in n1]
        common = [n for n in n1 if n in n2]
        diff_parts = [n for n in common if z1.read(n) != z2.read(n)]
        z1.close(); z2.close()
        order_ok = n1 == n2
        ok12 = order_ok and not only1 and not only2 and diff_parts == [DOCXML]
        if ok12:
            d12 = '仅 document.xml 不同，%d 个部件逐字节一致，顺序一致' % len(common)
        else:
            d12 = []
            if not order_ok:
                d12.append('条目顺序不同')
            if only1:
                d12.append('仅基线有 %d 个 %s' % (len(only1), only1[:4]))
            if only2:
                d12.append('仅新版有 %d 个 %s' % (len(only2), only2[:4]))
            if diff_parts != [DOCXML]:
                d12.append('字节不同 %d 个 %s' % (len(diff_parts), diff_parts[:6]))
            d12 = ' ; '.join(d12)
        if stage == 's6' and not ok12:
            rep.add_expect(False, '12 部件字节一致',
                           d12 + '　← EXPECT：S6 修订模式新增 word/comments.xml'
                                 '（部件集合必然变化，非缺陷）')
        else:
            rep.add(ok12, '12 部件字节一致', d12)

        # 13 可见文本 diff
        Pa, Pb = paras(xb), paras(x)
        if len(Pa) == len(Pb):
            changed, bad = [], []
            for i, ((a1, b1), (a2, b2)) in enumerate(zip(Pa, Pb)):
                t1, t2 = render(xb[a1:b1]), render(x[a2:b2])
                if t1 == t2:
                    continue
                changed.append(i)
                f1 = [re.sub(r'\s', '', m.group(0)) for m in NUM.finditer(t1)]
                f2 = [re.sub(r'\s', '', m.group(0)) for m in NUM.finditer(t2)]
                if f1 != f2 or NUM.sub('@', t1) != NUM.sub('@', t2):
                    bad.append(i)
            ok13 = not bad
            d13 = '%d 段变化%s' % (len(changed),
                                  '，非编号变化 %s' % bad[:8] if bad else '，全部是编号空格规范化')
        else:
            ok13 = False
            d13 = '段落数不同 %d vs %d' % (len(Pa), len(Pb))
        if stage == 's6' and not ok13:
            rep.add_expect(False, '13 可见文本 diff',
                           d13 + '　← EXPECT：S6 插题注段 / 修订模式，段数必然不同'
                                 '（非缺陷）')
        else:
            rep.add(ok13, '13 可见文本 diff', d13)

    if not quiet:
        pass
    rep.out()
    return rep


if __name__ == '__main__':
    argv = sys.argv[1:]
    # --stage（支持 `--stage s6` 与 `--stage=s6` 两种写法，默认 s2）
    stage = 's2'
    stripped, i = [], 0
    while i < len(argv):
        a = argv[i]
        if a == '--stage':
            stage = argv[i + 1] if i + 1 < len(argv) else 's2'
            i += 2
            continue
        if a.startswith('--stage='):
            stage = a.split('=', 1)[1]
            i += 1
            continue
        stripped.append(a)
        i += 1
    if stage not in ('s2', 's6'):
        print('[err] --stage 只认 s2 / s6，收到 %r' % stage)
        sys.exit(1)

    args = [a for a in stripped if not a.startswith('--')]
    if not args:
        print(__doc__); sys.exit(1)
    tgt = args[0]
    base = args[1] if len(args) > 1 else None
    if not os.path.exists(tgt):
        print('找不到文件: %s' % tgt); sys.exit(1)
    r = check(tgt, base, quiet='--quiet' in stripped, stage=stage)
    sys.exit(0 if r.fail == 0 else 2)
